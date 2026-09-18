"""Create and incrementally import the CRM and Cases Excel databases.

This module intentionally has no connection to the customer portal, dashboard,
or any mail automation.  It imports source columns exactly as they appear in
row 1 of each workbook.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sqlite3
from datetime import date, datetime, time
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


DATABASE_DIRECTORY = Path(__file__).resolve().parent.parent / "databases"
CRM_DATABASE_NAME = "crm_database.sqlite3"
CASES_DATABASE_NAME = "cases_database.sqlite3"


def quote(identifier: str) -> str:
    return '"' + identifier.replace('"', '""') + '"'


def cell_text(value: object) -> str | None:
    """Store source values losslessly enough for Excel data and keep IDs text."""
    if value is None:
        return None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def read_workbook(path: Path) -> tuple[str, list[str], Iterable[tuple[object, ...]]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    if len(workbook.worksheets) != 1:
        workbook.close()
        raise ValueError(f"{path.name} must contain exactly one data sheet; found {len(workbook.worksheets)}")
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as error:
        workbook.close()
        raise ValueError(f"{path.name} is empty") from error
    headers = [cell_text(value) for value in header_row]
    if any(not heading or not heading.strip() for heading in headers):
        workbook.close()
        raise ValueError(f"{path.name} has a blank column heading; columns cannot be safely preserved")
    if len(set(headers)) != len(headers):
        workbook.close()
        raise ValueError(f"{path.name} has duplicate column headings; columns cannot be safely preserved")
    return worksheet.title, headers, rows


def find_order_id_column(headers: list[str]) -> str:
    matches = [header for header in headers if header.strip().casefold() == "order id"]
    if len(matches) != 1:
        raise ValueError(f"Expected one exact 'Order ID' column, found {matches or 'none'}")
    return matches[0]


def row_hash(values: list[str | None]) -> str:
    encoded = json.dumps(values, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def initialise_database(connection: sqlite3.Connection, table_name: str, headers: list[str], order_id: str, is_crm: bool) -> None:
    existing = [row[1] for row in connection.execute(f"PRAGMA table_info({quote(table_name)})")]
    if not existing:
        columns = ", ".join(f"{quote(header)} TEXT" for header in headers)
        connection.execute(f"CREATE TABLE {quote(table_name)} ({columns})")
    elif existing != headers:
        raise ValueError(
            f"Existing {table_name} schema does not exactly match this workbook. "
            "No records were changed."
        )

    connection.execute(
        "CREATE TABLE IF NOT EXISTS imported_record_hashes "
        "(record_hash TEXT PRIMARY KEY, imported_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)"
    )
    index_name = "ux_crm_order_id" if is_crm else "ix_cases_order_id"
    unique = "UNIQUE " if is_crm else ""
    connection.execute(
        f"CREATE {unique}INDEX IF NOT EXISTS {quote(index_name)} "
        f"ON {quote(table_name)} ({quote(order_id)})"
    )


def import_workbook(source: Path, database_path: Path, table_name: str, is_crm: bool) -> dict[str, object]:
    sheet_name, headers, source_rows = read_workbook(source)
    order_id = find_order_id_column(headers)
    database_path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(database_path)
    connection.execute("PRAGMA journal_mode=WAL")
    connection.execute("PRAGMA synchronous=FULL")
    inserted = skipped = source_records = 0
    try:
        with connection:
            initialise_database(connection, table_name, headers, order_id, is_crm)
            known_hashes = {row[0] for row in connection.execute("SELECT record_hash FROM imported_record_hashes")}
            known_order_ids = (
                {row[0] for row in connection.execute(f"SELECT {quote(order_id)} FROM {quote(table_name)}") if row[0] is not None}
                if is_crm else set()
            )
            placeholders = ", ".join("?" for _ in headers)
            insert_sql = f"INSERT INTO {quote(table_name)} ({', '.join(quote(header) for header in headers)}) VALUES ({placeholders})"
            order_index = headers.index(order_id)
            for raw_values in source_rows:
                values = [cell_text(value) for value in raw_values]
                # XLSX rows commonly omit trailing empty cells. Pad those cells so
                # each value remains aligned with its original source heading.
                if len(values) > len(headers):
                    raise ValueError(f"Source row {source_records + 2} has more cells than the header row")
                values.extend([None] * (len(headers) - len(values)))
                if not any(value is not None for value in values):
                    continue
                source_records += 1
                digest = row_hash(values)
                # Cases use a complete-row hash for safe repeat-file imports. CRM
                # additionally enforces its permanent Order ID business key.
                if digest in known_hashes:
                    skipped += 1
                    continue
                if is_crm and values[order_index] is not None and values[order_index] in known_order_ids:
                    skipped += 1
                    continue
                try:
                    connection.execute(insert_sql, values)
                except sqlite3.IntegrityError:
                    if is_crm:
                        skipped += 1
                        continue
                    raise
                connection.execute("INSERT INTO imported_record_hashes(record_hash) VALUES (?)", (digest,))
                known_hashes.add(digest)
                if is_crm and values[order_index] is not None:
                    known_order_ids.add(values[order_index])
                inserted += 1
    finally:
        connection.close()
    return {
        "database": str(database_path), "table": table_name, "sheet": sheet_name,
        "columns": headers, "order_id_column": order_id, "source_records": source_records,
        "inserted": inserted, "skipped": skipped,
    }


def database_report(database_path: Path, table_name: str, order_id: str) -> dict[str, object]:
    connection = sqlite3.connect(database_path)
    try:
        total = connection.execute(f"SELECT COUNT(*) FROM {quote(table_name)}").fetchone()[0]
        duplicate_order_ids = connection.execute(
            f"SELECT COUNT(*) FROM (SELECT {quote(order_id)} FROM {quote(table_name)} "
            f"GROUP BY {quote(order_id)} HAVING COUNT(*) > 1)"
        ).fetchone()[0]
        indexes = connection.execute(f"PRAGMA index_list({quote(table_name)})").fetchall()
        return {"total_records": total, "duplicate_order_id_values": duplicate_order_ids, "indexes": indexes}
    finally:
        connection.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Incrementally import CRM and Cases .xlsx files into isolated SQLite databases.")
    parser.add_argument("--crm", type=Path, required=True)
    parser.add_argument("--cases", type=Path, required=True)
    arguments = parser.parse_args()

    crm = import_workbook(arguments.crm, DATABASE_DIRECTORY / CRM_DATABASE_NAME, "crm_records", True)
    cases = import_workbook(arguments.cases, DATABASE_DIRECTORY / CASES_DATABASE_NAME, "case_records", False)
    crm["verification"] = database_report(Path(crm["database"]), "crm_records", str(crm["order_id_column"]))
    cases["verification"] = database_report(Path(cases["database"]), "case_records", str(cases["order_id_column"]))
    print(json.dumps({"crm": crm, "cases": cases}, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
